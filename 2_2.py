"""
Один файл: SMS по признакам + статистика IMEI/IMSI из ClickHouse.

Установка:
    pip install clickhouse-connect python-docx

Рядом с программой нужен keywords.txt (один признак на строку).
Остальные файлы создаются автоматически.
"""

import argparse
import hashlib
import json
import logging
import os
import re
import sys
import time
from datetime import datetime, time as clock_time, timedelta
from pathlib import Path

if sys.version_info < (3, 10):
    raise SystemExit("Нужен Python 3.10 или новее.")

try:
    import clickhouse_connect
except ImportError as error:
    raise SystemExit(
        "Не найдена библиотека clickhouse-connect. "
        "Установка: python -m pip install clickhouse-connect"
    ) from error

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as error:
    raise SystemExit(
        "Не найдена библиотека python-docx. "
        "Установка: python -m pip install python-docx"
    ) from error


# =====================================================================
# ЗАПОЛНИ ЭТОТ БЛОК
# =====================================================================

CLICKHOUSE_HOST = "10.0.0.1"
CLICKHOUSE_PORT = 8123
CLICKHOUSE_USER = "default"
CLICKHOUSE_PASSWORD = "CHANGE_ME"
CLICKHOUSE_DATABASE = "default"
CLICKHOUSE_SECURE = False
CLICKHOUSE_VERIFY_TLS = True

SMS_TABLE = "VL_SMS"
SMS_DATETIME_COLUMN = "DateTime"
SMS_TEXT_COLUMN = "TpData"
SMS_PHONE_TO_COLUMN = "PhoneTo"
SMS_PHONE_FROM_COLUMN = "PhoneFrom"

LOCATION_TABLE = "WL_Location"
LOCATION_DATETIME_COLUMN = "DateTime"
LOCATION_IMEI_COLUMN = "IMEI"
LOCATION_IMSI_COLUMN = "IMSI"

# С этого момента строится новый общий отчёт.
FIRST_START = datetime(2026, 10, 27, 7, 0, 0)

# При первом запуске этот период становится базой уже известных IMEI, IMSI и пар.
# Чем больше период, тем точнее определяется "новый" идентификатор.
LOCATION_HISTORY_DAYS = 30

# Граница отчётного дня: 06:45 уже относится к следующему дню.
REPORT_BOUNDARY = clock_time(6, 45)

RUN_TIMES = [
    clock_time(3, 0),
    clock_time(6, 45),
    clock_time(7, 0),
    clock_time(11, 0),
    clock_time(15, 0),
    clock_time(19, 0),
    clock_time(23, 0),
]

OVERLAP_MINUTES = 24 * 60
RETRY_MINUTES = 5

# =====================================================================
# НИЖЕ НИЧЕГО МЕНЯТЬ НЕ НУЖНО
# =====================================================================

BASE_DIR = Path(__file__).resolve().parent
KEYWORDS_FILE = BASE_DIR / "keywords.txt"
SMS_HASHES_FILE = BASE_DIR / "combined_sms_hashes.txt"
IMEI_HASHES_FILE = BASE_DIR / "known_imei_hashes.txt"
IMSI_HASHES_FILE = BASE_DIR / "known_imsi_hashes.txt"
PAIR_HASHES_FILE = BASE_DIR / "known_pair_hashes.txt"
LOCATION_INITIALIZED_FILE = BASE_DIR / "location_initialized.txt"
CURSOR_FILE = BASE_DIR / "combined_last_query.txt"
STATE_FILE = BASE_DIR / "combined_report_state.json"
WORD_FILE = BASE_DIR / "common_report.docx"
IDENTIFIERS_FILE = BASE_DIR / "new_identifiers.txt"
LOG_FILE = BASE_DIR / "combined_monitor.log"
LOCK_FILE = BASE_DIR / "combined_monitor.lock"

IDENTIFIER_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def setup_logging():
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[logging.FileHandler(LOG_FILE, encoding="utf-8"), logging.StreamHandler()],
    )


def validate_settings():
    """Отлавливает ошибки в настройках до первого большого запроса."""
    if CLICKHOUSE_HOST in {"", "10.0.0.1"}:
        raise ValueError("Заполните CLICKHOUSE_HOST в начале файла.")
    if CLICKHOUSE_PASSWORD == "CHANGE_ME":
        raise ValueError("Заполните CLICKHOUSE_PASSWORD в начале файла.")
    if not isinstance(CLICKHOUSE_PORT, int) or not 1 <= CLICKHOUSE_PORT <= 65535:
        raise ValueError("CLICKHOUSE_PORT должен быть числом от 1 до 65535.")
    if LOCATION_HISTORY_DAYS < 1:
        raise ValueError("LOCATION_HISTORY_DAYS должно быть не меньше 1.")
    if not RUN_TIMES:
        raise ValueError("RUN_TIMES не должен быть пустым.")
    for value, description in [
        (CLICKHOUSE_DATABASE, "имя базы"),
        (SMS_TABLE, "SMS-таблица"),
        (SMS_DATETIME_COLUMN, "SMS DateTime"),
        (SMS_TEXT_COLUMN, "TpData"),
        (SMS_PHONE_TO_COLUMN, "PhoneTo"),
        (SMS_PHONE_FROM_COLUMN, "PhoneFrom"),
        (LOCATION_TABLE, "Location-таблица"),
        (LOCATION_DATETIME_COLUMN, "Location DateTime"),
        (LOCATION_IMEI_COLUMN, "IMEI"),
        (LOCATION_IMSI_COLUMN, "IMSI"),
    ]:
        checked_identifier(value, description)
    load_keywords()


def acquire_lock():
    """Не даёт запустить две копии программы одновременно."""
    try:
        descriptor = os.open(LOCK_FILE, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as error:
        raise RuntimeError(
            f"Уже существует {LOCK_FILE.name}. "
            "Если программа точно не запущена, удалите этот lock-файл."
        ) from error
    os.write(descriptor, str(os.getpid()).encode("ascii"))
    os.close(descriptor)
    return LOCK_FILE


def checked_identifier(value, description):
    if not IDENTIFIER_RE.fullmatch(value):
        raise ValueError(f"Неверное {description}: {value!r}")
    return f"`{value}`"


def sha256(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def imei_hash(imei):
    return sha256(imei)


def imsi_hash(imsi):
    return sha256(imsi)


def pair_hash(imei, imsi):
    # Разделитель не даёт склейкам "12+345" и "123+45" дать один результат.
    return sha256(imei + "\x1f" + imsi)


def load_hashes(filename):
    if not filename.exists():
        return set()
    with filename.open("r", encoding="ascii", errors="ignore") as file:
        return {line.strip() for line in file if line.strip()}


def append_hashes(filename, values):
    if not values:
        return
    with filename.open("a", encoding="ascii", newline="\n") as file:
        for value in sorted(values):
            file.write(value + "\n")
        file.flush()
        os.fsync(file.fileno())


def write_hashes(filename, values):
    temporary = filename.with_suffix(filename.suffix + ".tmp")
    with temporary.open("w", encoding="ascii", newline="\n") as file:
        for value in sorted(values):
            file.write(value + "\n")
        file.flush()
        os.fsync(file.fileno())
    temporary.replace(filename)


def load_keywords():
    if not KEYWORDS_FILE.exists():
        raise FileNotFoundError(f"Не найден {KEYWORDS_FILE}")
    with KEYWORDS_FILE.open("r", encoding="utf-8-sig") as file:
        result = [line.strip() for line in file if line.strip() and not line.lstrip().startswith("#")]
    if not result:
        raise ValueError("keywords.txt пуст")
    return result


def find_first_keyword(text, keywords):
    for keyword in keywords:
        if re.search(rf"(?<!\w){re.escape(keyword)}(?!\w)", text, flags=re.IGNORECASE):
            return keyword
    return None


def get_client():
    return clickhouse_connect.get_client(
        host=CLICKHOUSE_HOST,
        port=CLICKHOUSE_PORT,
        username=CLICKHOUSE_USER,
        password=CLICKHOUSE_PASSWORD,
        database=CLICKHOUSE_DATABASE,
        secure=CLICKHOUSE_SECURE,
        verify=CLICKHOUSE_VERIFY_TLS,
        connect_timeout=15,
        send_receive_timeout=180,
    )


def read_table_columns(client, table_name):
    query = """
        SELECT name, type
        FROM system.columns
        WHERE database = {database:String}
          AND table = {table:String}
    """
    result = client.query(
        query,
        parameters={"database": CLICKHOUSE_DATABASE, "table": table_name},
    )
    return {str(row[0]): str(row[1]) for row in result.result_rows}


def validate_database(client):
    """Проверяет таблицы, колонки и тип DateTime до запуска цикла."""
    client.command("SELECT 1")
    checks = [
        (
            SMS_TABLE,
            [SMS_DATETIME_COLUMN, SMS_TEXT_COLUMN, SMS_PHONE_TO_COLUMN, SMS_PHONE_FROM_COLUMN],
            SMS_DATETIME_COLUMN,
        ),
        (
            LOCATION_TABLE,
            [LOCATION_DATETIME_COLUMN, LOCATION_IMEI_COLUMN, LOCATION_IMSI_COLUMN],
            LOCATION_DATETIME_COLUMN,
        ),
    ]
    for table_name, required, datetime_column in checks:
        columns = read_table_columns(client, table_name)
        if not columns:
            raise RuntimeError(
                f"Таблица {CLICKHOUSE_DATABASE}.{table_name} не найдена "
                "или у учётной записи нет прав на её просмотр."
            )
        missing = [name for name in required if name not in columns]
        if missing:
            raise RuntimeError(
                f"В {table_name} не найдены колонки: {', '.join(missing)}. "
                f"Доступны: {', '.join(sorted(columns))}"
            )
        datetime_type = columns[datetime_column]
        if not (datetime_type.startswith("DateTime") or datetime_type.startswith("Nullable(DateTime")):
            raise RuntimeError(
                f"Колонка {table_name}.{datetime_column} имеет тип {datetime_type}, "
                "а ожидался DateTime или DateTime64."
            )
        logging.info("Проверена таблица %s.%s", CLICKHOUSE_DATABASE, table_name)


def query_sms(client, start, end, keywords):
    database = checked_identifier(CLICKHOUSE_DATABASE, "имя базы")
    table = checked_identifier(SMS_TABLE, "SMS-таблица")
    dt = checked_identifier(SMS_DATETIME_COLUMN, "SMS DateTime")
    text = checked_identifier(SMS_TEXT_COLUMN, "TpData")
    phone_to = checked_identifier(SMS_PHONE_TO_COLUMN, "PhoneTo")
    phone_from = checked_identifier(SMS_PHONE_FROM_COLUMN, "PhoneFrom")

    parameters = {
        "start": start.strftime("%Y-%m-%d %H:%M:%S"),
        "end": end.strftime("%Y-%m-%d %H:%M:%S"),
    }
    conditions = []
    for number, keyword in enumerate(keywords):
        name = f"keyword_{number}"
        conditions.append(f"positionCaseInsensitiveUTF8(toString({text}), {{{name}:String}}) > 0")
        parameters[name] = keyword

    query = f"""
        SELECT
            toString({dt}) AS DateTime,
            toString({phone_from}) AS PhoneFrom,
            toString({phone_to}) AS PhoneTo,
            toString({text}) AS TpData
        FROM {database}.{table}
        WHERE {dt} >= toDateTime({{start:String}})
          AND {dt} < toDateTime({{end:String}})
          AND ({' OR '.join(conditions)})
        ORDER BY {dt}
    """
    return list(client.query(query, parameters=parameters).named_results())


def query_location(client, start, end):
    database = checked_identifier(CLICKHOUSE_DATABASE, "имя базы")
    table = checked_identifier(LOCATION_TABLE, "Location-таблица")
    dt = checked_identifier(LOCATION_DATETIME_COLUMN, "Location DateTime")
    imei = checked_identifier(LOCATION_IMEI_COLUMN, "IMEI")
    imsi = checked_identifier(LOCATION_IMSI_COLUMN, "IMSI")
    parameters = {
        "start": start.strftime("%Y-%m-%d %H:%M:%S"),
        "end": end.strftime("%Y-%m-%d %H:%M:%S"),
    }
    query = f"""
        SELECT
            toString({imei}) AS IMEI,
            toString({imsi}) AS IMSI,
            toString(min({dt})) AS FirstSeen
        FROM {database}.{table}
        WHERE {dt} >= toDateTime({{start:String}})
          AND {dt} < toDateTime({{end:String}})
          AND notEmpty(toString({imei}))
          AND notEmpty(toString({imsi}))
        GROUP BY {imei}, {imsi}
        ORDER BY FirstSeen
    """
    return list(client.query(query, parameters=parameters).named_results())


def report_date_for(value):
    if value.time() >= REPORT_BOUNDARY:
        return value.date().isoformat()
    return (value.date() - timedelta(days=1)).isoformat()


def parse_datetime(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    text = str(value).strip()
    if text.endswith("Z"):
        text = text[:-1] + "+00:00"
    return datetime.fromisoformat(text).replace(tzinfo=None)


def load_cursor():
    if not CURSOR_FILE.exists():
        return FIRST_START
    try:
        saved = datetime.fromisoformat(CURSOR_FILE.read_text(encoding="ascii").strip())
        return max(FIRST_START, saved - timedelta(minutes=OVERLAP_MINUTES))
    except (OSError, ValueError):
        return FIRST_START


def save_cursor(value):
    temporary = CURSOR_FILE.with_suffix(".tmp")
    temporary.write_text(value.isoformat(timespec="seconds"), encoding="ascii")
    temporary.replace(CURSOR_FILE)


def empty_state():
    return {"days": {}}


def load_state():
    if not STATE_FILE.exists():
        return empty_state()
    try:
        with STATE_FILE.open("r", encoding="utf-8-sig") as file:
            state = json.load(file)
    except (OSError, json.JSONDecodeError) as error:
        raise RuntimeError(
            f"Не удалось прочитать {STATE_FILE.name}. "
            "Не удаляйте файл; сначала сделайте его копию."
        ) from error
    state.setdefault("days", {})
    return state


def save_state(state):
    temporary = STATE_FILE.with_suffix(".json.tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as file:
        json.dump(state, file, ensure_ascii=False, indent=2)
        file.flush()
        os.fsync(file.fileno())
    temporary.replace(STATE_FILE)


def get_day(state, day):
    value = state["days"].setdefault(
        day,
        {"sms": {}, "new_imei": {}, "new_imsi": {}, "new_pairs": {}},
    )
    value.setdefault("sms", {})
    value.setdefault("new_imei", {})
    value.setdefault("new_imsi", {})
    value.setdefault("new_pairs", {})
    return value


def initialize_location_history(client):
    if LOCATION_INITIALIZED_FILE.exists():
        return
    start = FIRST_START - timedelta(days=LOCATION_HISTORY_DAYS)
    logging.info("Первичная загрузка WL_Location: %s - %s", start, FIRST_START)
    rows = query_location(client, start, FIRST_START)
    imei_values = set()
    imsi_values = set()
    pair_values = set()
    for row in rows:
        imei = str(row.get("IMEI", "")).strip()
        imsi = str(row.get("IMSI", "")).strip()
        if not imei or not imsi:
            continue
        imei_values.add(imei_hash(imei))
        imsi_values.add(imsi_hash(imsi))
        pair_values.add(pair_hash(imei, imsi))
    write_hashes(IMEI_HASHES_FILE, imei_values)
    write_hashes(IMSI_HASHES_FILE, imsi_values)
    write_hashes(PAIR_HASHES_FILE, pair_values)
    LOCATION_INITIALIZED_FILE.write_text(
        f"Initialized {datetime.now().isoformat(timespec='seconds')}\n"
        f"History {start} - {FIRST_START}\n",
        encoding="utf-8",
    )
    logging.info(
        "Известно из истории: IMEI=%s, IMSI=%s, пар=%s",
        len(imei_values), len(imsi_values), len(pair_values),
    )


def process_sms(rows, keywords, state):
    known = load_hashes(SMS_HASHES_FILE)
    added = set()
    for row in rows:
        text = str(row.get("TpData") or "")
        keyword = find_first_keyword(text, keywords) if text else None
        if not keyword:
            continue
        value = sha256(text)
        if value in known or value in added:
            continue
        event_time = parse_datetime(row["DateTime"])
        day = get_day(state, report_date_for(event_time))
        day["sms"].setdefault(
            value,
            {
                "keyword": keyword,
                "PhoneFrom": str(row.get("PhoneFrom") or ""),
                "PhoneTo": str(row.get("PhoneTo") or ""),
                "DateTime": event_time.strftime("%Y-%m-%d %H:%M:%S"),
                "TpData": text,
            },
        )
        added.add(value)
    return added


def process_location(rows, state):
    known_imei = load_hashes(IMEI_HASHES_FILE)
    known_imsi = load_hashes(IMSI_HASHES_FILE)
    known_pairs = load_hashes(PAIR_HASHES_FILE)
    add_imei = set()
    add_imsi = set()
    add_pairs = set()

    normalized = []
    for row in rows:
        imei = str(row.get("IMEI") or "").strip()
        imsi = str(row.get("IMSI") or "").strip()
        if not imei or not imsi:
            continue
        first_seen = parse_datetime(row["FirstSeen"])
        normalized.append((first_seen, imei, imsi))
    normalized.sort()

    for first_seen, imei, imsi in normalized:
        ih = imei_hash(imei)
        sh = imsi_hash(imsi)
        ph = pair_hash(imei, imsi)
        day = get_day(state, report_date_for(first_seen))
        common = {
            "IMEI": imei,
            "IMSI": imsi,
            "FirstSeen": first_seen.strftime("%Y-%m-%d %H:%M:%S"),
        }
        if ih not in known_imei and ih not in add_imei:
            day["new_imei"].setdefault(ih, common)
            add_imei.add(ih)
        if sh not in known_imsi and sh not in add_imsi:
            day["new_imsi"].setdefault(sh, common)
            add_imsi.add(sh)
        if ph not in known_pairs and ph not in add_pairs:
            day["new_pairs"].setdefault(ph, common)
            add_pairs.add(ph)

    return add_imei, add_imsi, add_pairs


def set_cell(cell, value, bold=False, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


def add_heading(document, text, level):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_sms_table(document, records):
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Признак", "Отправитель", "Получатель", "Дата", "Текст SMS"]
    repeat_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell(cell, header, bold=True)
        shade_cell(cell, "D9EAF7")
    for number, record in enumerate(sorted(records, key=lambda x: (x["keyword"].casefold(), x["DateTime"]))):
        cells = table.add_row().cells
        values = [record["keyword"], record["PhoneFrom"], record["PhoneTo"], record["DateTime"], record["TpData"]]
        for cell, value in zip(cells, values):
            set_cell(cell, value)
            if number % 2:
                shade_cell(cell, "F4F8FB")


def build_word(state):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = section.bottom_margin = Cm(1.3)
    section.left_margin = section.right_margin = Cm(1.3)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Ежедневный отчёт по SMS и устройствам")
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(0, 0, 0)

    days = sorted(state["days"])
    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run(f"Периодов в отчёте: {len(days)}")

    for index, day_text in enumerate(days):
        if index:
            document.add_page_break()
        day = get_day(state, day_text)
        add_heading(document, f"Данные за {day_text}", 1)

        sms = list(day["sms"].values())
        add_heading(document, f"SMS по признакам: {len(sms)}", 2)
        if sms:
            add_sms_table(document, sms)
        else:
            document.add_paragraph("Новых SMS по признакам нет.")

        add_heading(document, "Статистика WL Location", 2)
        stats = document.add_table(rows=1, cols=2)
        stats.style = "Table Grid"
        set_cell(stats.rows[0].cells[0], "Показатель", bold=True)
        set_cell(stats.rows[0].cells[1], "Количество", bold=True)
        shade_cell(stats.rows[0].cells[0], "D9EAF7")
        shade_cell(stats.rows[0].cells[1], "D9EAF7")
        rows = [
            ("Новые IMEI", len(day["new_imei"])),
            ("Новые IMSI", len(day["new_imsi"])),
            ("Новые связки IMEI + IMSI", len(day["new_pairs"])),
        ]
        for label, value in rows:
            cells = stats.add_row().cells
            set_cell(cells[0], label)
            set_cell(cells[1], value)

        document.add_paragraph("Сами IMEI и IMSI записаны в new_identifiers.txt.")

    temporary = WORD_FILE.with_name("." + WORD_FILE.name + ".tmp.docx")
    document.save(temporary)
    temporary.replace(WORD_FILE)


def build_identifiers_text(state):
    temporary = IDENTIFIERS_FILE.with_suffix(".txt.tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as file:
        for day_text in sorted(state["days"]):
            day = get_day(state, day_text)
            file.write(f"===== {day_text} =====\n\n")

            file.write(f"НОВЫЕ IMEI: {len(day['new_imei'])}\n")
            file.write("IMEI\tIMSI\tFIRST_SEEN\n")
            for row in sorted(day["new_imei"].values(), key=lambda x: x["FirstSeen"]):
                file.write(f"{row['IMEI']}\t{row['IMSI']}\t{row['FirstSeen']}\n")

            file.write(f"\nНОВЫЕ IMSI: {len(day['new_imsi'])}\n")
            file.write("IMSI\tIMEI\tFIRST_SEEN\n")
            for row in sorted(day["new_imsi"].values(), key=lambda x: x["FirstSeen"]):
                file.write(f"{row['IMSI']}\t{row['IMEI']}\t{row['FirstSeen']}\n")

            file.write(f"\nНОВЫЕ СВЯЗКИ IMEI + IMSI: {len(day['new_pairs'])}\n")
            file.write("IMEI\tIMSI\tFIRST_SEEN\n")
            for row in sorted(day["new_pairs"].values(), key=lambda x: x["FirstSeen"]):
                file.write(f"{row['IMEI']}\t{row['IMSI']}\t{row['FirstSeen']}\n")
            file.write("\n")
        file.flush()
        os.fsync(file.fileno())
    temporary.replace(IDENTIFIERS_FILE)


def run_cycle(client):
    initialize_location_history(client)
    keywords = load_keywords()
    start = load_cursor()
    end = datetime.now().replace(microsecond=0)
    if end <= FIRST_START:
        logging.info("FIRST_START ещё не наступило")
        return

    logging.info("Запрос за период %s - %s", start, end)
    sms_rows = query_sms(client, start, end, keywords)
    location_rows = query_location(client, start, end)
    logging.info("SMS-%s, Location-пар: %s", len(sms_rows), len(location_rows))

    state = load_state()
    new_sms_hashes = process_sms(sms_rows, keywords, state)
    new_imei_hashes, new_imsi_hashes, new_pair_hashes = process_location(location_rows, state)

    # Сначала сохраняем состояние и отчёты, и только потом двигаем курсор.
    save_state(state)
    build_word(state)
    build_identifiers_text(state)
    append_hashes(SMS_HASHES_FILE, new_sms_hashes)
    append_hashes(IMEI_HASHES_FILE, new_imei_hashes)
    append_hashes(IMSI_HASHES_FILE, new_imsi_hashes)
    append_hashes(PAIR_HASHES_FILE, new_pair_hashes)
    save_cursor(end)

    logging.info(
        "Добавлено: SMS=%s, IMEI=%s, IMSI=%s, пар=%s",
        len(new_sms_hashes), len(new_imei_hashes), len(new_imsi_hashes), len(new_pair_hashes),
    )


def next_run_after(now):
    candidates = []
    for offset in (0, 1):
        candidate_date = now.date() + timedelta(days=offset)
        for scheduled_time in RUN_TIMES:
            candidate = datetime.combine(candidate_date, scheduled_time)
            if candidate > now:
                candidates.append(candidate)
    return min(candidates)


def main():
    parser = argparse.ArgumentParser(description="SMS и IMEI/IMSI из ClickHouse")
    parser.add_argument(
        "--once",
        action="store_true",
        help="выполнить один цикл и завершиться",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="только проверить настройки, подключение и схему БД",
    )
    args = parser.parse_args()
    setup_logging()
    lock_file = acquire_lock()
    client = None
    try:
        validate_settings()
        logging.info("Запуск объединённого монитора")
        client = get_client()
        validate_database(client)
        if args.check:
            logging.info("Проверка завершена успешно. Данные не изменялись.")
            return
        if args.once:
            run_cycle(client)
            logging.info("Однократный запуск завершён.")
            return
        while True:
            try:
                run_cycle(client)
                next_run = next_run_after(datetime.now())
                logging.info("Следующий запуск: %s", next_run)
                time.sleep(max(1, (next_run - datetime.now()).total_seconds()))
            except KeyboardInterrupt:
                raise
            except Exception:
                logging.exception("Ошибка. Повтор через %s минут", RETRY_MINUTES)
                time.sleep(RETRY_MINUTES * 60)
                try:
                    client.close()
                except Exception:
                    pass
                client = get_client()
    except KeyboardInterrupt:
        logging.info("Программа остановлена")
    finally:
        if client is not None:
            try:
                client.close()
            except Exception:
                pass
        try:
            lock_file.unlink()
        except FileNotFoundError:
            pass


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        logging.exception("Критическая ошибка запуска")
        print(f"\nОШИБКА: {error}")
        print(f"Подробности записаны в: {LOG_FILE}")
        raise SystemExit(1)
